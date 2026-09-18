import io
from datetime import datetime, date, time
from typing import List, Dict, Any, Tuple
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from sqlalchemy.orm import Session
from sqlalchemy import select

from app.models.employee import Employee
from app.models.attendance import Attendance
from app.models.enums import CheckMethod


def export_attendance_to_excel(attendance_records: List[Attendance]) -> io.BytesIO:
    """
    Generate an Excel sheet (.xlsx) of attendance records.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Report"

    # Header styling
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    headers = [
        "Employee Code",
        "Employee Name",
        "Date",
        "Check In Time",
        "Check Out Time",
        "Total Hours",
        "Method",
        "Status",
        "Regularized"
    ]
    ws.append(headers)

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Rows
    for att in attendance_records:
        emp_code = att.employee.employee_code if att.employee else ""
        emp_name = att.employee.full_name if att.employee else ""
        date_str = att.date.strftime("%Y-%m-%d") if att.date else ""
        in_str = att.check_in_time.strftime("%H:%M:%S") if att.check_in_time else ""
        out_str = att.check_out_time.strftime("%H:%M:%S") if att.check_out_time else ""
        hours = float(att.total_hours) if att.total_hours is not None else 0.0
        method = att.check_in_method.value if att.check_in_method else ""
        status = att.status or "present"
        reg = "Yes" if att.is_regularized else "No"

        ws.append([
            emp_code,
            emp_name,
            date_str,
            in_str,
            out_str,
            hours,
            method,
            status,
            reg
        ])

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def import_attendance_from_excel(
    file_bytes: bytes,
    db: Session
) -> Tuple[int, int, List[Dict[str, Any]]]:
    """
    Import attendance from uploaded .xlsx file.
    Expected columns: Employee Code, Date (YYYY-MM-DD), Check In (HH:MM or HH:MM:SS), Check Out (HH:MM or HH:MM:SS)
    Returns (imported_count, skipped_count, errors_list)
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active

    imported = 0
    skipped = 0
    errors = []

    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        return 0, 0, [{"row": 1, "reason": "Empty or missing data sheet"}]

    # Header is row 1, data starts at row 2
    for row_idx, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue  # empty row

        emp_code = str(row[0]).strip() if row[0] is not None else ""
        raw_date = row[1]
        raw_check_in = row[2] if len(row) > 2 else None
        raw_check_out = row[3] if len(row) > 3 else None

        if not emp_code:
            errors.append({"row": row_idx, "reason": "Missing employee code"})
            skipped += 1
            continue

        employee = db.execute(select(Employee).where(Employee.employee_code == emp_code)).scalar_one_or_none()
        if not employee:
            errors.append({"row": row_idx, "reason": f"Employee code '{emp_code}' not found"})
            skipped += 1
            continue

        # Parse date
        parsed_date = None
        if isinstance(raw_date, (datetime, date)):
            parsed_date = raw_date if isinstance(raw_date, date) else raw_date.date()
        elif isinstance(raw_date, str):
            for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
                try:
                    parsed_date = datetime.strptime(raw_date.strip(), fmt).date()
                    break
                except ValueError:
                    pass

        if not parsed_date:
            errors.append({"row": row_idx, "reason": "Invalid date format"})
            skipped += 1
            continue

        # Parse times
        def parse_time(raw_val, base_d: date) -> Optional[datetime]:
            if not raw_val:
                return None
            if isinstance(raw_val, datetime):
                return raw_val
            if isinstance(raw_val, time):
                return datetime.combine(base_d, raw_val)
            if isinstance(raw_val, str):
                for t_fmt in ("%H:%M:%S", "%H:%M", "%I:%M %p"):
                    try:
                        t = datetime.strptime(raw_val.strip(), t_fmt).time()
                        return datetime.combine(base_d, t)
                    except ValueError:
                        pass
            return None

        dt_check_in = parse_time(raw_check_in, parsed_date)
        dt_check_out = parse_time(raw_check_out, parsed_date)

        total_hours = None
        if dt_check_in and dt_check_out:
            total_hours = round((dt_check_out - dt_check_in).total_seconds() / 3600.0, 2)

        # Check existing attendance
        att = db.execute(
            select(Attendance).where(
                Attendance.employee_id == employee.id,
                Attendance.date == parsed_date
            )
        ).scalar_one_or_none()

        if att:
            # Update existing
            att.check_in_time = dt_check_in or att.check_in_time
            att.check_out_time = dt_check_out or att.check_out_time
            if total_hours is not None:
                att.total_hours = total_hours
            att.check_in_method = CheckMethod.manual_override
        else:
            # Create new
            att = Attendance(
                employee_id=employee.id,
                date=parsed_date,
                check_in_time=dt_check_in,
                check_out_time=dt_check_out,
                total_hours=total_hours,
                check_in_method=CheckMethod.manual_override,
                check_out_method=CheckMethod.manual_override if dt_check_out else None,
                status="present"
            )
            db.add(att)

        imported += 1

    db.commit()
    return imported, skipped, errors


def export_leaves_to_excel(leave_requests: List[Any]) -> io.BytesIO:
    """
    Generate an Excel sheet of leave requests.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leave Requests"

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    headers = [
        "Request ID",
        "Employee Code",
        "Employee Name",
        "Leave Type",
        "Start Date",
        "End Date",
        "Half Day",
        "Status",
        "Applied At",
        "Decided At",
        "Reason"
    ]
    ws.append(headers)

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for lr in leave_requests:
        emp_code = lr.employee.employee_code if lr.employee else ""
        emp_name = lr.employee.full_name if lr.employee else ""
        lt_name = lr.leave_type.name if lr.leave_type else ""
        start_str = lr.start_date.strftime("%Y-%m-%d") if lr.start_date else ""
        end_str = lr.end_date.strftime("%Y-%m-%d") if lr.end_date else ""
        half = "Yes" if lr.is_half_day else "No"
        app_str = lr.applied_at.strftime("%Y-%m-%d %H:%M") if lr.applied_at else ""
        dec_str = lr.decided_at.strftime("%Y-%m-%d %H:%M") if lr.decided_at else ""

        ws.append([
            lr.id,
            emp_code,
            emp_name,
            lt_name,
            start_str,
            end_str,
            half,
            lr.status,
            app_str,
            dec_str,
            lr.reason or ""
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def import_leaves_from_excel(
    file_bytes: bytes,
    db: Session
) -> Tuple[int, int, List[Dict[str, Any]]]:
    """
    Bulk import/adjust leave balances from Excel.
    Columns: Employee Code, Leave Type, Year, Total Allotted
    """
    from app.models.leave import LeaveType, LeaveBalance
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active

    imported = 0
    skipped = 0
    errors = []

    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        return 0, 0, [{"row": 1, "reason": "Empty or missing data sheet"}]

    for row_idx, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue

        emp_code = str(row[0]).strip() if row[0] is not None else ""
        lt_name = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        raw_year = row[2] if len(row) > 2 else None
        raw_allotted = row[3] if len(row) > 3 else None

        if not emp_code:
            errors.append({"row": row_idx, "reason": "Missing employee code"})
            skipped += 1
            continue

        employee = db.execute(select(Employee).where(Employee.employee_code == emp_code)).scalar_one_or_none()
        if not employee:
            errors.append({"row": row_idx, "reason": f"Employee code '{emp_code}' not found"})
            skipped += 1
            continue

        leave_type = db.execute(select(LeaveType).where(LeaveType.name.ilike(lt_name))).scalar_one_or_none()
        if not leave_type:
            errors.append({"row": row_idx, "reason": f"Leave type '{lt_name}' not found"})
            skipped += 1
            continue

        try:
            year = int(raw_year)
            allotted = float(raw_allotted)
        except (TypeError, ValueError):
            errors.append({"row": row_idx, "reason": "Invalid year or allotted days format"})
            skipped += 1
            continue

        bal = db.execute(
            select(LeaveBalance).where(
                LeaveBalance.employee_id == employee.id,
                LeaveBalance.leave_type_id == leave_type.id,
                LeaveBalance.year == year
            )
        ).scalar_one_or_none()

        if bal:
            diff = allotted - float(bal.total_allotted)
            bal.total_allotted = allotted
            bal.balance = float(bal.balance) + diff
        else:
            bal = LeaveBalance(
                employee_id=employee.id,
                leave_type_id=leave_type.id,
                year=year,
                total_allotted=allotted,
                used=0.0,
                balance=allotted
            )
            db.add(bal)

        imported += 1

    db.commit()
    return imported, skipped, errors
